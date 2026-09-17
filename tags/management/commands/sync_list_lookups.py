from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from tags.models import Category, IssueType, RedTag, Section, Station


def _as_text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


class Command(BaseCommand):
    help = (
        'Rebuild Section/Station and Category/IssueType master data from the '
        'Excel List sheet (Section/Station_Alias + Defects Classification).'
    )

    def add_arguments(self, parser):
        parser.add_argument('xlsx_path', type=str, help='Path to Red Tag Report .xlsx')

    @transaction.atomic
    def handle(self, *args, **options):
        path = Path(options['xlsx_path'])
        if not path.exists():
            raise CommandError(f'File not found: {path}')

        wb = load_workbook(path, data_only=True, read_only=True)
        if 'List' not in wb.sheetnames:
            raise CommandError('Workbook has no List sheet')

        ws = wb['List']
        header = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
        rows = list(ws.iter_rows(min_row=3, values_only=True))

        master_sections = set()
        section_station_pairs = []
        for row in rows:
            section_code = _as_text(row[15] if len(row) > 15 else None)
            station_name = _as_text(row[16] if len(row) > 16 else None)
            if section_code and station_name:
                master_sections.add(section_code)
                section_station_pairs.append((section_code, station_name))

        # Deduplicate while preserving order
        seen = set()
        unique_pairs = []
        for pair in section_station_pairs:
            key = (pair[0].upper(), pair[1].upper())
            if key in seen:
                continue
            seen.add(key)
            unique_pairs.append(pair)

        section_by_code = {}
        for code in sorted(master_sections, key=str.upper):
            section, _ = Section.objects.update_or_create(
                code=code,
                defaults={'show_in_form': True},
            )
            section_by_code[code.upper()] = section

        # Hide sections that are not in the List sheet master table
        Section.objects.exclude(
            code__in=list(master_sections)
        ).update(show_in_form=False)

        station_by_key = {}
        master_station_ids = []
        for section_code, station_name in unique_pairs:
            section = section_by_code[section_code.upper()]
            station, _ = Station.objects.update_or_create(
                section=section,
                name=station_name,
                defaults={'show_in_form': True},
            )
            station_by_key[(section.code.upper(), station_name.upper())] = station
            master_station_ids.append(station.id)

        Station.objects.exclude(id__in=master_station_ids).update(show_in_form=False)

        remapped = 0
        created_for_history = 0
        for tag in RedTag.objects.select_related('section', 'station').iterator():
            if not tag.section_id or not tag.station_id:
                continue
            key = (tag.section.code.upper(), tag.station.name.upper())
            target = station_by_key.get(key)
            if not target:
                # Keep historical station under the tag's section if missing from List
                target, was_created = Station.objects.get_or_create(
                    section=tag.section,
                    name=tag.station.name,
                )
                station_by_key[key] = target
                created_for_history += int(was_created)
            if tag.station_id != target.id:
                tag.station = target
                tag.save(update_fields=['station'])
                remapped += 1

        # Remove stations that are unused and not part of the List master pairs
        keep_ids = {s.id for s in station_by_key.values()}
        deleted_orphans, _ = (
            Station.objects.exclude(id__in=keep_ids)
            .filter(red_tags__isnull=True)
            .delete()
        )
        # Defects Classification: categories in row 2 cols J-N, issue types in rows below
        category_headers = [_as_text(h) for h in header[9:14]]
        if not any(category_headers):
            category_headers = [
                'PART HANDLING', 'ASSEMBLY MMO', 'PROCESS DEVIATION', 'KDQR', 'PAINT DEFECTS'
            ]

        issue_count = 0
        for col_idx, cat_name in enumerate(category_headers):
            if not cat_name:
                continue
            category, _ = Category.objects.get_or_create(name=cat_name)
            for row in rows:
                issue_name = _as_text(row[9 + col_idx] if len(row) > 9 + col_idx else None)
                if not issue_name:
                    continue
                _, created = IssueType.objects.get_or_create(
                    category=category, name=issue_name
                )
                issue_count += int(created)

        self.stdout.write(self.style.SUCCESS(
            f'Master sections={len(master_sections)}, '
            f'section/station pairs={len(unique_pairs)}, '
            f'remapped tags={remapped}, '
            f'history stations created={created_for_history}, '
            f'orphan stations removed={deleted_orphans}, '
            f'new issue types={issue_count}'
        ))
        self.stdout.write(
            'Form sections will be those with stations from the List sheet Section/Station_Alias table.'
        )
