import datetime
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from tags.models import (
    Category, Dealer, Employee, IssueType, RedTag, Section, Station,
    Vehicle, VehicleModel,
)


def _as_text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _as_date(value):
    if value is None or value == '':
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    return None


def _as_int(value):
    if value is None or value == '':
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def _normalize_status(value):
    raw = _as_text(value).upper()
    if raw in ('C', 'CLOSED'):
        return RedTag.STATUS_CLOSED
    return RedTag.STATUS_PENDING


def _normalize_location(value):
    raw = _as_text(value).upper()
    if raw in ('EOL', 'FINAL'):
        return raw
    if raw:
        return 'OTHER'
    return ''


class Command(BaseCommand):
    help = 'Import Red Tag Report Excel workbook (staff, lookups, and red tag rows).'

    def add_arguments(self, parser):
        parser.add_argument('xlsx_path', type=str, help='Path to Red Tag Report .xlsx')
        parser.add_argument(
            '--limit',
            type=int,
            default=0,
            help='Optional max number of Red Tag rows to import (0 = all)',
        )
        parser.add_argument(
            '--skip-tags',
            action='store_true',
            help='Only import lookup sheets (Staff / List / Models)',
        )
        parser.add_argument(
            '--clear-tags',
            action='store_true',
            help='Delete existing red tags before import',
        )

    def handle(self, *args, **options):
        path = Path(options['xlsx_path'])
        if not path.exists():
            raise CommandError(f'File not found: {path}')

        self.stdout.write(f'Loading {path} ...')
        wb = load_workbook(path, data_only=True, read_only=True)

        self._import_staff(wb)
        self._import_list_lookups(wb)
        self._import_models(wb)

        if options['clear_tags']:
            deleted, _ = RedTag.objects.all().delete()
            self.stdout.write(self.style.WARNING(f'Deleted {deleted} existing red tags'))

        if not options['skip_tags']:
            self._import_red_tags(wb, limit=options['limit'])

        self.stdout.write(self.style.SUCCESS('Import complete.'))

    def _import_staff(self, wb):
        if 'Staff' not in wb.sheetnames:
            self.stdout.write('Staff sheet missing — skipped')
            return

        ws = wb['Staff']
        created = updated = 0
        for i, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
            emp_no = _as_text(row[1] if len(row) > 1 else None)
            name = _as_text(row[2] if len(row) > 2 else None)
            if not emp_no or not name:
                continue
            _, was_created = Employee.objects.update_or_create(
                employee_no=emp_no,
                defaults={'full_name': name, 'active': True},
            )
            created += int(was_created)
            updated += int(not was_created)
        self.stdout.write(f'Staff: {created} created, {updated} updated')

    def _import_list_lookups(self, wb):
        if 'List' not in wb.sheetnames:
            self.stdout.write('List sheet missing — skipped')
            return

        ws = wb['List']
        header_row = list(ws.iter_rows(min_row=2, max_row=2, values_only=True))[0]
        rows = list(ws.iter_rows(min_row=3, values_only=True))

        # Section / Station_Alias master table (columns P/Q)
        for row in rows:
            section_code = _as_text(row[15] if len(row) > 15 else None)
            station_name = _as_text(row[16] if len(row) > 16 else None)
            if not section_code or not station_name:
                continue
            section, _ = Section.objects.get_or_create(
                code=section_code,
                defaults={'show_in_form': True},
            )
            if not section.show_in_form:
                section.show_in_form = True
                section.save(update_fields=['show_in_form'])
            Station.objects.update_or_create(
                section=section,
                name=station_name,
                defaults={'show_in_form': True},
            )

        # Defects Classification: categories in row 2 cols J-N, issue types below
        category_headers = [_as_text(h) for h in header_row[9:14]]
        if not any(category_headers):
            category_headers = [
                'PART HANDLING', 'ASSEMBLY MMO', 'PROCESS DEVIATION', 'KDQR', 'PAINT DEFECTS'
            ]

        for col_idx, cat_name in enumerate(category_headers):
            if not cat_name:
                continue
            category, _ = Category.objects.get_or_create(name=cat_name)
            for row in rows:
                issue_name = _as_text(row[9 + col_idx] if len(row) > 9 + col_idx else None)
                if issue_name:
                    IssueType.objects.get_or_create(category=category, name=issue_name)

        # Dealers / models from List columns F-H
        for row in rows:
            model_code = _as_text(row[5] if len(row) > 5 else None)
            model_name = _as_text(row[6] if len(row) > 6 else None)
            dealer_name = _as_text(row[7] if len(row) > 7 else None)
            dealer = None
            if dealer_name and dealer_name != '#VALUE!':
                dealer, _ = Dealer.objects.get_or_create(name=dealer_name)
            if model_name:
                VehicleModel.objects.update_or_create(
                    name=model_name,
                    defaults={'code': model_code, 'dealer': dealer},
                )

        self.stdout.write(
            f'Lookups: sections={Section.objects.filter(stations__isnull=False).distinct().count()}, '
            f'categories={Category.objects.count()}, '
            f'issue_types={IssueType.objects.count()}, '
            f'stations={Station.objects.count()}, '
            f'models={VehicleModel.objects.count()}, '
            f'dealers={Dealer.objects.count()}'
        )

    def _import_models(self, wb):
        if 'Models' not in wb.sheetnames:
            return
        ws = wb['Models']
        header = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
        dealers = []
        for name in header:
            text = _as_text(name)
            if text:
                dealer, _ = Dealer.objects.get_or_create(name=text)
                dealers.append(dealer)
            else:
                dealers.append(None)

        for row in ws.iter_rows(min_row=3, values_only=True):
            for idx, model_name in enumerate(row):
                text = _as_text(model_name)
                if not text or idx >= len(dealers) or not dealers[idx]:
                    continue
                VehicleModel.objects.update_or_create(
                    name=text,
                    defaults={'dealer': dealers[idx]},
                )

    def _import_red_tags(self, wb, limit=0):
        sheet_name = 'Red Tag' if 'Red Tag' in wb.sheetnames else wb.sheetnames[0]
        ws = wb[sheet_name]

        # Warm caches
        employees = {e.employee_no: e for e in Employee.objects.all()}
        sections = {s.code: s for s in Section.objects.all()}
        categories = {c.name.upper(): c for c in Category.objects.all()}
        issue_types = {
            (it.category_id, it.name.upper()): it for it in IssueType.objects.select_related('category')
        }
        stations = {
            (s.section_id, s.name.upper()): s
            for s in Station.objects.select_related('section')
            if s.section_id
        }        models = {m.name.upper(): m for m in VehicleModel.objects.all()}
        dealers = {d.name.upper(): d for d in Dealer.objects.all()}
        vehicles = {v.chassis_no: v for v in Vehicle.objects.select_related('model')}

        created = skipped = 0
        batch = []
        batch_size = 1000

        for i, row in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
            if limit and created >= limit:
                break

            date_raised = _as_date(row[0] if len(row) > 0 else None)
            chassis_no = _as_text(row[5] if len(row) > 5 else None)
            issue_description = _as_text(row[9] if len(row) > 9 else None)

            if not date_raised or not chassis_no or not issue_description:
                skipped += 1
                continue

            section_code = _as_text(row[1] if len(row) > 1 else None) or 'UNKNOWN'
            dealer_name = _as_text(row[2] if len(row) > 2 else None)
            model_name = _as_text(row[3] if len(row) > 3 else None) or 'UNKNOWN'
            lot = _as_text(row[4] if len(row) > 4 else None)
            part_number = _as_text(row[6] if len(row) > 6 else None)
            part_name = _as_text(row[7] if len(row) > 7 else None)
            quantity = _as_int(row[8] if len(row) > 8 else None)
            category_name = _as_text(row[10] if len(row) > 10 else None)
            issue_type_name = _as_text(row[11] if len(row) > 11 else None)
            reason_code = _as_text(row[12] if len(row) > 12 else None)
            raised_by_no = _as_text(row[13] if len(row) > 13 else None)
            station_name = _as_text(row[14] if len(row) > 14 else None)
            action_taken = _as_text(row[15] if len(row) > 15 else None)
            status = _normalize_status(row[16] if len(row) > 16 else None)
            closing_date = _as_date(row[17] if len(row) > 17 else None)
            verified_by_no = _as_text(row[18] if len(row) > 18 else None)
            remarks = _as_text(row[19] if len(row) > 19 else None)
            location = _normalize_location(row[20] if len(row) > 20 else None)
            vi_size = _as_text(row[22] if len(row) > 22 else None)

            if part_number.upper() == 'N/A':
                part_number = ''
            if part_name.upper() == 'N/A':
                part_name = ''
            if dealer_name.upper() in ('#VALUE!', 'NONE'):
                dealer_name = ''

            section = sections.get(section_code)
            if not section:
                section = Section.objects.create(code=section_code)
                sections[section_code] = section

            dealer = None
            if dealer_name:
                dealer = dealers.get(dealer_name.upper())
                if not dealer:
                    dealer = Dealer.objects.create(name=dealer_name)
                    dealers[dealer_name.upper()] = dealer

            model = models.get(model_name.upper())
            if not model:
                model = VehicleModel.objects.create(name=model_name, dealer=dealer)
                models[model_name.upper()] = model
            elif dealer and not model.dealer_id:
                model.dealer = dealer
                model.save(update_fields=['dealer'])

            vehicle = vehicles.get(chassis_no)
            if not vehicle:
                vehicle = Vehicle.objects.create(
                    chassis_no=chassis_no,
                    model=model,
                    dealer=dealer,
                    lot=lot,
                    vi_size=vi_size,
                )
                vehicles[chassis_no] = vehicle
            else:
                changed = False
                if lot and vehicle.lot != lot:
                    vehicle.lot = lot
                    changed = True
                if dealer and vehicle.dealer_id != dealer.id:
                    vehicle.dealer = dealer
                    changed = True
                if vi_size and vehicle.vi_size != vi_size:
                    vehicle.vi_size = vi_size
                    changed = True
                if changed:
                    vehicle.save()

            category = None
            if category_name:
                category = categories.get(category_name.upper())
                if not category:
                    category = Category.objects.create(name=category_name)
                    categories[category_name.upper()] = category

            issue_type = None
            if category and issue_type_name:
                key = (category.id, issue_type_name.upper())
                issue_type = issue_types.get(key)
                if not issue_type:
                    issue_type = IssueType.objects.create(category=category, name=issue_type_name)
                    issue_types[key] = issue_type

            station = None
            if station_name:
                station_key = (section.id, station_name.upper())
                station = stations.get(station_key)
                if not station:
                    station = Station.objects.create(section=section, name=station_name)
                    stations[station_key] = station

            raised_by = employees.get(raised_by_no) if raised_by_no else None
            if raised_by_no and not raised_by:
                raised_by = Employee.objects.create(
                    employee_no=raised_by_no,
                    full_name=f'Employee {raised_by_no}',
                )
                employees[raised_by_no] = raised_by

            verified_by = employees.get(verified_by_no) if verified_by_no else None
            if verified_by_no and not verified_by:
                verified_by = Employee.objects.create(
                    employee_no=verified_by_no,
                    full_name=f'Employee {verified_by_no}',
                )
                employees[verified_by_no] = verified_by

            batch.append(RedTag(
                date_raised=date_raised,
                vehicle=vehicle,
                section=section,
                station=station,
                part_number=part_number,
                part_name=part_name,
                quantity=quantity,
                issue_description=issue_description,
                category=category,
                issue_type=issue_type,
                reason_code=reason_code,
                raised_by=raised_by,
                verified_by=verified_by,
                action_taken=action_taken,
                status=status,
                closing_date=closing_date,
                remarks=remarks,
                location=location,
            ))
            created += 1

            if len(batch) >= batch_size:
                with transaction.atomic():
                    RedTag.objects.bulk_create(batch, batch_size=batch_size)
                self.stdout.write(f'  imported {created} tags...')
                batch = []

        if batch:
            with transaction.atomic():
                RedTag.objects.bulk_create(batch, batch_size=batch_size)

        self.stdout.write(self.style.SUCCESS(
            f'Red tags imported: {created} (skipped incomplete rows: {skipped})'
        ))
